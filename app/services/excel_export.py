from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from sqlalchemy import extract, func
from sqlalchemy.orm import Session

from app.models import Budget, Category, Movement, SavingsGoal


def export_monthly_report(db: Session, year: int, month: int, output_dir: str = "exports"):
    out_path = Path(output_dir)
    out_path.mkdir(parents=True, exist_ok=True)
    report_file = out_path / f"reporte_{year}_{month:02}.xlsx"

    wb = Workbook()
    ws_mov = wb.active
    ws_mov.title = "Movimientos"
    ws_mov.append(
        ["Fecha", "Tipo", "Monto", "Metodo", "Categoria", "Nota"]
    )

    movements = (
        db.query(Movement, Category.name)
        .join(Category, Category.id == Movement.category_id)
        .filter(
            extract("year", Movement.movement_date) == year,
            extract("month", Movement.movement_date) == month,
        )
        .order_by(Movement.movement_date.asc())
        .all()
    )

    for movement, category_name in movements:
        ws_mov.append(
            [
                movement.movement_date.isoformat(),
                movement.movement_type,
                movement.amount,
                movement.payment_method,
                category_name,
                movement.note or "",
            ]
        )

    ws_res = wb.create_sheet("Resumen")
    ws_res.append(["Indicador", "Valor"])

    income = sum(m.amount for m, _ in movements if m.movement_type == "income")
    expense = sum(m.amount for m, _ in movements if m.movement_type == "expense")
    balance = income - expense

    ws_res.append(["Ingresos", income])
    ws_res.append(["Gastos", expense])
    ws_res.append(["Balance", balance])

    chart = BarChart()
    chart.title = "Ingresos vs Gastos"
    data = Reference(ws_res, min_col=2, min_row=1, max_row=3)
    cats = Reference(ws_res, min_col=1, min_row=2, max_row=3)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws_res.add_chart(chart, "D2")

    ws_cat = wb.create_sheet("Categorias")
    ws_cat.append(["Categoria", "Gasto"])
    expense_by_category = (
        db.query(Category.name, func.sum(Movement.amount))
        .join(Movement, Movement.category_id == Category.id)
        .filter(
            extract("year", Movement.movement_date) == year,
            extract("month", Movement.movement_date) == month,
            Movement.movement_type == "expense",
        )
        .group_by(Category.name)
        .all()
    )

    for name, total in expense_by_category:
        ws_cat.append([name, float(total or 0)])

    if len(expense_by_category) > 0:
        pie = PieChart()
        pie.title = "Distribucion de gastos por categoria"
        labels = Reference(ws_cat, min_col=1, min_row=2, max_row=len(expense_by_category) + 1)
        data = Reference(ws_cat, min_col=2, min_row=1, max_row=len(expense_by_category) + 1)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        ws_cat.add_chart(pie, "D2")

    ws_budget = wb.create_sheet("Presupuestos")
    ws_budget.append(["Categoria", "Limite", "Gasto", "Diferencia"])
    budgets = (
        db.query(Budget, Category.name)
        .join(Category, Category.id == Budget.category_id)
        .filter(Budget.year == year, Budget.month == month)
        .all()
    )
    expense_lookup = {name: float(total or 0) for name, total in expense_by_category}
    for budget, category_name in budgets:
        spent = expense_lookup.get(category_name, 0.0)
        ws_budget.append([category_name, budget.limit_amount, spent, budget.limit_amount - spent])

    ws_goal = wb.create_sheet("Ahorro")
    ws_goal.append(["Meta", "Objetivo", "Actual", "Falta"])
    goals = db.query(SavingsGoal).order_by(SavingsGoal.name.asc()).all()
    for goal in goals:
        ws_goal.append(
            [
                goal.name,
                goal.target_amount,
                goal.current_amount,
                max(goal.target_amount - goal.current_amount, 0),
            ]
        )

    wb.save(report_file)
    return str(report_file)
